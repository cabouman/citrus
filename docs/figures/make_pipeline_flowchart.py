"""Build docs/figures/citrus_pipeline.pptx, the flowchart of the scatter-corrected MBIR pipeline of CITRUS (one
16:9 slide, editable in PowerPoint), and export it to docs/source/_static/citrus_pipeline.png with PowerPoint.

    python docs/figures/make_pipeline_flowchart.py            # writes the .pptx and, on Windows with PowerPoint, the .png

Edit the slide in PowerPoint and re-export (File > Save As > PNG, or rerun this script) when the pipeline changes.
"""
from __future__ import annotations

import os
import subprocess
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
PPTX = os.path.join(HERE, "citrus_pipeline.pptx")
PNG = os.path.join(HERE, "..", "source", "_static", "citrus_pipeline.png")

FONT = "Calibri"
# fill, border, text colours per role
STYLE = {
    "data": ("DEEAF6", "2F5597", "1F3864"),      # measured data, initial and final volumes (blue)
    "density": ("E2F0D9", "548235", "385723"),   # density model and grids (green)
    "scatter": ("FBE5D6", "C55A11", "843C0C"),   # LBTE scatter solver (orange)
    "primary": ("E4DFEC", "7030A0", "3F1F5E"),   # primary (purple)
    "correct": ("FFF2CC", "BF9000", "7F6000"),   # ratio and correction (yellow)
    "mbir": ("DEEAF6", "2F5597", "1F3864"),      # MBIR (blue)
    "note": ("F2F2F2", "7F7F7F", "404040"),      # notes (grey)
}
EMU_PER_INCH = 914400


def rgb(hex6: str) -> RGBColor:
    return RGBColor.from_string(hex6)


def add_box(slide, x, y, w, h, text, role, *, size=11, bold_first=True, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
            dashed=False, fill=True):
    fill_c, line_c, text_c = STYLE[role]
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.12
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill_c)
    else:
        s.fill.background()
    s.line.color.rgb = rgb(line_c)
    s.line.width = Pt(1.25)
    if dashed:
        ln = s.line._get_or_add_ln()
        ln.append(ln.makeelement(qn("a:prstDash"), {"val": "dash"}))
    s.shadow.inherit = False
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = line
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bool(bold_first and i == 0)
        r.font.color.rgb = rgb(text_c)
    return s


def add_label(slide, x, y, w, h, text, color="404040", size=10, bold=False, align=PP_ALIGN.LEFT, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0.0)
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = rgb(color)
    return tb


def _arrow_end(line_elm, kind="triangle"):
    ln = line_elm._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": kind, "w": "med", "len": "med"}))


def add_arrow(slide, x1, y1, x2, y2, color="595959", width=1.5, head=True, dashed=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    if dashed:
        ln = c.line._get_or_add_ln()
        ln.append(ln.makeelement(qn("a:prstDash"), {"val": "dash"}))
    if head:
        _arrow_end(c.line)
    return c


def add_polyline_arrow(slide, points, color="595959", width=1.5, dashed=False):
    """Axis-aligned segments through `points` [(x, y), ...] in inches; the last segment carries the arrowhead."""
    for i in range(len(points) - 1):
        (x1, y1), (x2, y2) = points[i], points[i + 1]
        add_arrow(slide, x1, y1, x2, y2, color=color, width=width, head=(i == len(points) - 2), dashed=dashed)


def left(s):
    return s.left / EMU_PER_INCH


def top(s):
    return s.top / EMU_PER_INCH


def right(s):
    return Emu(s.left + s.width).inches


def bottom(s):
    return Emu(s.top + s.height).inches


def cx(s):
    return Emu(s.left + s.width // 2).inches


def cy(s):
    return Emu(s.top + s.height // 2).inches


def build() -> str:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    L = 3.05  # left edge of the loop container

    # ---- inputs (left column) -----------------------------------------------------------------------------------
    y_data = add_box(slide, 0.35, 0.95, 2.25, 0.95,
                     "Measured projections\ncounts Y = primary + scatter,\nflat field I0, view angles", "data")
    fdk = add_box(slide, 0.35, 2.45, 2.25, 0.95,
                  "FDK initialization\nParker weights for a short scan\n(MBIRTORCH)", "mbir")
    x0 = add_box(slide, 0.35, 3.95, 2.25, 0.7, "Initial volume X₀", "data", size=11)
    add_arrow(slide, cx(y_data), bottom(y_data), cx(fdk), top(fdk))
    add_arrow(slide, cx(fdk), bottom(fdk), cx(x0), top(x0))

    # ---- outer loop container ------------------------------------------------------------------------------------
    loop = add_box(slide, L, 0.45, 13.0 - L, 5.95, "", "note", shape=MSO_SHAPE.ROUNDED_RECTANGLE, dashed=True,
                   fill=False)
    loop.adjustments[0] = 0.04
    loop.line.color.rgb = rgb("595959")
    add_label(slide, L + 0.25, 0.5, 9.5, 0.3,
              "Outer iteration k:  one scatter step per view  +  one MBIR iteration   (repeated N times)",
              color="404040", size=11, bold=True)

    # row A: density model and grids
    dens = add_box(slide, 3.5, 1.0, 3.9, 1.05,
                   "Density model\nXₖ → densities of the anchor materials\n"
                   "(two materials per voxel, one map at a time)", "density")
    grids = add_box(slide, 7.8, 1.0, 5.0, 1.05,
                    "Transport and quadrature grids\nconservative resampling onto cubic transport cells (1 cm)\n"
                    "and q³ quadrature points per cell, any voxel size", "density")
    add_arrow(slide, right(dens), cy(dens), left(grids), cy(grids))

    # row B: the LBTE scatter solver (per view)
    band_left = 3.95
    solver_bg = add_box(slide, band_left, 2.35, 12.85 - band_left, 1.95, "", "scatter", shape=MSO_SHAPE.RECTANGLE,
                        fill=True)
    solver_bg.fill.fore_color.rgb = rgb("FDF3EC")
    solver_bg.line.color.rgb = rgb("C55A11")
    solver_bg.line.width = Pt(0.75)
    add_label(slide, band_left + 0.1, 2.38, 8.5, 0.28,
              "LBTE scatter solver, one step per view v (deterministic, multigroup, discrete ordinates)",
              color="843C0C", size=10, bold=True)
    uncol = add_box(slide, 4.1, 2.75, 2.75, 1.25,
                    "Uncollided fluence\nray traced through the quadrature grid,\nfirst-collision source on its points",
                    "scatter")
    sweep = add_box(slide, 7.15, 2.75, 2.95, 1.25,
                    "One transport sweep\nlinear-discontinuous cells of C, fed by\nthe first-collision source + the\n"
                    "aged accumulator Φᵥ⁽ᵏ⁻¹⁾ → Φᵥ⁽ᵏ⁾",
                    "scatter", size=10)
    proj = add_box(slide, 10.4, 2.75, 2.35, 1.25,
                   "Collapsed-cone projection\nfirst-order + multiple-collision\nemission to the panel\n"
                   "→ scatter sᵥ", "scatter", size=10)
    add_arrow(slide, right(uncol), cy(uncol), left(sweep), cy(sweep))
    add_arrow(slide, right(sweep), cy(sweep), left(proj), cy(proj))
    add_polyline_arrow(slide, [(cx(sweep) + 0.9, bottom(sweep)), (cx(sweep) + 0.9, bottom(sweep) + 0.17),
                               (cx(sweep) - 0.9, bottom(sweep) + 0.17), (cx(sweep) - 0.9, bottom(sweep) + 0.01)],
                       color="C55A11", width=1.25, dashed=True)
    add_label(slide, cx(sweep) - 1.4, bottom(sweep) + 0.02, 2.8, 0.2,
              "Φᵥ kept per view across outer iterations (age-ordered source iteration)",
              color="843C0C", size=8, italic=True, align=PP_ALIGN.CENTER)

    # row C: primary, ratio, correction, MBIR
    prim = add_box(slide, 3.5, 4.75, 2.3, 1.05,
                   "Primary pᵥ of every view\nMBIRTORCH per-material projections,\nspectrum + detector response",
                   "primary", size=10)
    ratio = add_box(slide, 6.05, 4.75, 2.2, 1.05,
                    "Scatter-to-primary ratio\nSPRᵥ = sᵥ / pᵥ\n(interpolated in angle for\nskipped views)",
                    "correct", size=10)
    corr = add_box(slide, 8.5, 4.75, 2.0, 1.05,
                   "Corrected counts\nYᵥ / (1 + SPRᵥ)\nline integrals\n−log(Yᵥ / I0)", "correct",
                   size=10)
    mbir = add_box(slide, 10.75, 4.75, 2.0, 1.05,
                   "One warm-started\nMBIR iteration\nMBIRTORCH coordinate descent,\n"
                   "resumed error sinogram → Xₖ₊₁", "mbir", size=10)
    add_arrow(slide, right(prim), cy(prim), left(ratio), cy(ratio))
    add_arrow(slide, right(ratio), cy(ratio), left(corr), cy(corr))
    add_arrow(slide, right(corr), cy(corr), left(mbir), cy(mbir))

    # vertical flow: density -> uncollided, grids -> sweep, density -> primary (left margin), projection -> ratio
    add_arrow(slide, cx(uncol), bottom(dens), cx(uncol), top(uncol))
    add_arrow(slide, cx(sweep), bottom(grids), cx(sweep), top(solver_bg))
    add_polyline_arrow(slide, [(3.7, bottom(dens)), (3.7, top(prim))], color="548235", width=1.25)
    add_label(slide, 3.28, 3.2, 0.5, 0.2, "ρₘ", color="385723", size=9, italic=True)
    add_polyline_arrow(slide, [(cx(proj), bottom(proj)), (cx(proj), 4.52), (cx(ratio) + 0.3, 4.52),
                               (cx(ratio) + 0.3, top(ratio))])
    add_label(slide, cx(ratio) + 0.38, 4.5, 0.6, 0.2, "sᵥ", color="843C0C", size=9, italic=True)

    # X0 into the loop; X_{k+1} back to the density model
    add_polyline_arrow(slide, [(right(x0), cy(x0)), (2.85, cy(x0)), (2.85, cy(dens) - 0.2),
                               (left(dens), cy(dens) - 0.2)])
    add_polyline_arrow(slide, [(right(mbir), cy(mbir)), (12.9, cy(mbir)), (12.9, 6.15), (3.25, 6.15),
                               (3.25, cy(dens) + 0.2), (left(dens), cy(dens) + 0.2)], color="2F5597", width=1.5)
    add_label(slide, 6.5, 5.95, 3.5, 0.2, "Xₖ₊₁ → next outer iteration", color="1F3864", size=9,
              italic=True, align=PP_ALIGN.CENTER)

    # output
    out = add_box(slide, 10.75, 6.65, 2.0, 0.65, "Scatter-corrected\nreconstruction X", "data", size=10)
    add_arrow(slide, cx(mbir), bottom(mbir), cx(out), top(out))

    # options note and caption
    add_box(slide, 0.35, 5.0, 2.25, 1.35,
            "Options\nscatter stage every n outer iterations;\nscatter step on every n-th view, the\n"
            "others interpolated in angle; anchor\nmaterials chosen from the catalogue", "note", size=9)
    add_label(slide, 0.35, 6.75, 8.5, 0.45,
              "CITRUS: Coupled Iterative Tomographic reconstruction and Recursively-Updated Scatter.\n"
              "With the densities held fixed, N steps reproduce N source iterations; the scatter converges with the "
              "reconstruction.", color="404040", size=9, italic=True)

    prs.save(PPTX)
    return PPTX


def export_png(pptx_path: str, png_path: str, width_px: int = 2400) -> bool:
    """Export slide 1 to PNG with PowerPoint (Windows COM via PowerShell). Returns True on success."""
    if sys.platform != "win32":
        return False
    os.makedirs(os.path.dirname(png_path), exist_ok=True)
    ps = (
        "$app = New-Object -ComObject PowerPoint.Application; "
        f"$p = $app.Presentations.Open('{os.path.abspath(pptx_path)}', $true, $true, $false); "
        f"$p.Slides.Item(1).Export('{os.path.abspath(png_path)}', 'PNG', {width_px}, {int(width_px * 7.5 / 13.333)}); "
        "$p.Close(); $app.Quit()"
    )
    r = subprocess.run(["powershell", "-NoProfile", "-Command", ps], capture_output=True, text=True)
    if r.returncode != 0:
        print(r.stdout, r.stderr, file=sys.stderr)
    return r.returncode == 0 and os.path.exists(png_path)


if __name__ == "__main__":
    path = build()
    print("wrote", path)
    if export_png(path, PNG):
        print("exported", os.path.abspath(PNG))
    else:
        print("PNG export skipped or failed (needs PowerPoint on Windows)")
